# -*- coding: utf-8 -*-
"""Universal table Excel export tool.

Background: Excel generation capabilities were once tied to the valuation tool (generate_valuation_excel_tool).
That is a data assembly  → LLM Generate specs (max 3 attempts)→ Rendering s heavy-duty pipeline
 (heavy Level 180 seconds, actual measurement 2-6 minutes). The user puts the model that has been calculated in the conversation
 (Sales/revenue forecast, etc.) was incorrectly routed to the valuation tool when exporting to Excel and the model was regenerated.
Resulting in a 420 second timeout.

The division of labor between this tool and valuation tools:
- This tool:**Deterministic export**——LLM Transmit the Markdown table generated in the conversation as it is.
  Pure openpyxl rendering, zero LLM calls, completed in seconds. Any put the data/table just now
  All your needs for making it into Excel should go here.
- Valuation tool: data from local annual reports**Self-built**Valuation model (profit deduction formula + PE scenario),
  Only if explicitly requested by the user"Valuation Calculation Sheet/Valuation Model"when used.

Parameter design (Markdown transparent transmission instead of structured JSON):
LLM There is already a ready-made Markdown table in the context (shown in the reply), you can just paste the original text.
Structured JSON parameters (nested sheets/columns/rows) instead provide LLM with contrivance/rewriting
digital opportunities——Transparent transmission of the original text is correct"Data cannot be tampered with"The observance of ironclad rules.
"""

import logging
import re
import time
from datetime import datetime
from pathlib import Path
from typing import Annotated, List, Optional, Tuple

from langchain_core.tools import tool
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from core.tools.base import register_tool

logger = logging.getLogger(__name__)

_OUTPUT_DIR = "data/deliverables"

# Markdown Table separated rows (|---|---:| etc.)
_TABLE_SEP_RE = re.compile(r"^\s*\|?[\s:|-]+\|?\s*$")
# cell text → numerical value
_INT_RE = re.compile(r"^[+-]?\d+$")
_THOUSANDS_RE = re.compile(r"^[+-]?\d{1,3}(,\d{3})+(\.\d+)?$")
_NUMBER_RE = re.compile(r"^[+-]?\d+(\.\d+)?$")
_PERCENT_RE = re.compile(r"^([+-]?\d+(?:\.\d+)?)\s*%$")


def _parse_cell(text: str) -> Tuple[object, Optional[str]]:
    """cell text → (value, Number format). Returned when it cannot be parsed as a numeric value (original text, None). """
    s = (text or "").strip()
    if not s:
        return "", None
    # Thousandths:1,234 / 12,345.67
    if _THOUSANDS_RE.match(s):
        return float(s.replace(",", "")), "#,##0.##"
    if _INT_RE.match(s):
        return int(s), "#,##0"
    m = _PERCENT_RE.match(s)
    if m:
        return float(m.group(1)) / 100.0, "0.00%"
    if _NUMBER_RE.match(s):
        return float(s), "0.##"
    return s, None


def _split_row(line: str) -> List[str]:
    """Markdown table row → List of cells (tolerating any variation of leading and trailing vertical bars)."""
    s = line.strip()
    if s.startswith("|"):
        s = s[1:]
    if s.endswith("|"):
        s = s[:-1]
    return [c.strip() for c in s.split("|")]


def _is_table_row(line: str) -> bool:
    return "|" in line and not _TABLE_SEP_RE.match(line)


def _parse_markdown_tables(
    markdown: str,
) -> List[Tuple[str, List[List[str]]]]:
    """Parse Markdown text → [(sheet name, Table row list), ...]. 

    For multiple sheets `## Title` Separate; if there is no title, all go into a single sheet.
    Each paragraph allows multiple consecutive tables (the beginning and the end are automatically spliced into one, and the number of columns is based on the table header).
    """
    # press ## title section
    segments: List[Tuple[str, List[str]]] = []
    current_title = ""
    current_lines: List[str] = []
    for line in (markdown or "").splitlines():
        m = re.match(r"^#{1,3}\s+(.+)$", line.strip())
        if m:
            if any(_is_table_row(l) for l in current_lines):
                segments.append((current_title, list(current_lines)))
            current_title = m.group(1).strip()
            current_lines = []
        else:
            current_lines.append(line)
    if any(_is_table_row(l) for l in current_lines):
        segments.append((current_title, list(current_lines)))

    if not segments:
        return []

    result: List[Tuple[str, List[List[str]]]] = []
    for title, lines in segments:
        rows: List[List[str]] = []
        for line in lines:
            if _is_table_row(line):
                rows.append(_split_row(line))
        if rows:
            result.append((title or f"表格{len(result) + 1}", rows))
    return result


def _display_width(text: str) -> float:
    """Estimated display width: 2 for Chinese/full-width characters."""
    return sum(2 if ord(ch) > 0x2E80 else 1 for ch in str(text))


def _sanitize_sheet_name(name: str, used: set) -> str:
    """Clean sheet names (Excel limitations:≤31 characters, disabled :\\/?*[]) . """
    clean = re.sub(r"[:\\/?*\[\]]", "_", (name or "").strip())[:31] or "Sheet"
    base, i = clean, 2
    while clean in used:
        suffix = f"_{i}"
        clean = base[: 31 - len(suffix)] + suffix
        i += 1
    used.add(clean)
    return clean


def _sanitize_filename_part(title: str) -> str:
    """Title → File name security fragment (retain Chinese characters, remove illegal characters).

    half-angle brackets () Must be replaced with full-width: unescaped half-width brackets in markdown link target
    The URL will be truncated according to CommonMark rules, resulting in incomplete download links.
    """
    slug = re.sub(r"[\\/:*?\"<>|\s]+", "_", (title or "").strip())
    slug = slug.replace("(", "（").replace(")", "）")
    slug = re.sub(r"_{2,}", "_", slug).strip("_")
    return slug[:40] or "table"


@tool
@register_tool(
    tool_id="export_table_excel_tool",
    name="表格Excel导出",
    description=(
        "Export the data table (original Markdown table) generated in the conversation to an Excel file (xlsx),"
        "Supports multiple tables divided into Sheets. Pure deterministic rendering, no data is regenerated or modified, and it is completed in seconds."
        "All numbers to be exported have already appeared in the conversation (forecast results, calculation tables, comparison tables, etc.),"
        "No matter what the user says'made Excel''Generate data table''Make a model table'Still'Export', This tool should be called"
        "Remodeling or switching to other modeling tools for this purpose is strictly prohibited."
    ),
    category="utility",
    is_online=False,
    auto_register=True,
    timeout_tier="light",
    capability_tags=[
        "excel_export", "table_export", "deliverable_generation",
        "markdown_to_excel", "data_export",
    ],
    tool_role_hint="specialized",
    output_shape="structured_markdown",
    when_to_use=(
        "The data/table/prediction results to be exported have been generated in the current conversation (the Markdown table has appeared in the LLM reply),"
        "Called when the user now requires them to be made into Excel/datasheet files——Just transmit the original text of the form verbatim."
        "For multiple tables '## table name' Separate. The user s statement may be'Make a model''Generate data table''Export'Wait,"
        "Use this tool whenever the data comes from conversational context and does not require recalculation."
    ),
    when_not_to_use=(
        "User explicit request'Valuation calculation sheet/valuation model/profit deduction/PE scenario' (Requires formula linkage"
        "financial model), use generate_valuation_excel_tool."
    ),
    returns=(
        "Markdown Text: export results, summary of rows and columns of each Sheet, file download link;"
        "Return an explicit failure description when there is no parsable table."
    ),
    example=(
        "export_table_excel_tool(title='Cyrus revenue forecast', "
        "markdown_tables='## Monthly Sales Forecast\\n| month | M9 | M7 |\\n|---|---:|---:|\\n| 2026-09 | 10500 | 8000 |')"
    ),
    related_tools=["generate_valuation_excel_tool"],
)
def export_table_excel_tool(
    title: Annotated[str, "Table theme (for file names and document titles), such as'Cyrus revenue forecast'"],
    markdown_tables: Annotated[
        str,
        "Markdown The original text of the form (directly pass through the form generated in the conversation, copy it verbatim, and it is strictly prohibited to rewrite any numbers)."
        "For multiple forms '## table name' Header rows are separated, and each header becomes a Sheet in Excel.",
    ],
    notes: Annotated[str, "Optional: Additional explanation (data source, caliber, assumptions, etc.), written at the bottom of the table"] = "",
) -> str:
    """Export the Markdown table in the conversation to an Excel file.

    Args:
        title: Table theme (file name)
        markdown_tables: Markdown Original form text, used for multiple sheets ## title separator
        notes: Optional footnote

    Returns:
        Markdown text(result summary + Download link, or failure instructions)
    """
    t0 = time.time()
    try:
        parsed = _parse_markdown_tables(markdown_tables or "")
        if not parsed:
            return (
                "### Table export failed\n\n"
                "The Markdown table was not parsed in the incoming content (required `|` Separated header rows and data rows).\n\n"
                "*Please confirm the table contents to be exported and try again.*"
            )

        wb = Workbook()
        wb.remove(wb.active)

        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill("solid", fgColor="4472C4")
        header_align = Alignment(horizontal="center", vertical="center")
        used_names: set = set()
        sheet_summaries: List[str] = []

        for sheet_title, rows in parsed:
            ws = wb.create_sheet(title=_sanitize_sheet_name(sheet_title, used_names))
            # Header = First line (the following delimited lines have been skipped during parsing)
            header = rows[0]
            data_rows = rows[1:]
            # Filter empty lines
            data_rows = [r for r in data_rows if any(c for c in r)]

            for col_idx, cell_text in enumerate(header, start=1):
                cell = ws.cell(row=1, column=col_idx, value=cell_text)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = header_align

            col_widths = [_display_width(h) + 4 for h in header]

            for row_idx, row in enumerate(data_rows, start=2):
                for col_idx in range(len(header)):
                    text = row[col_idx] if col_idx < len(row) else ""
                    value, num_fmt = _parse_cell(text)
                    cell = ws.cell(row=row_idx, column=col_idx + 1, value=value)
                    if num_fmt:
                        cell.number_format = num_fmt
                        cell.alignment = Alignment(horizontal="right")
                    if col_idx < len(row):
                        w = _display_width(text) + 2
                        if w > col_widths[col_idx]:
                            col_widths[col_idx] = w

            for col_idx, width in enumerate(col_widths, start=1):
                ws.column_dimensions[get_column_letter(col_idx)].width = min(width, 50)

            # Footnotes:notes + AI Generate Statement (Compliance Requirements)
            foot_row = len(data_rows) + 3
            footnotes = []
            if notes:
                footnotes.append(f"Description:{notes}")
            footnotes.append("This table is generated by AI and is for research reference only and does not constitute investment advice.")
            for i, note in enumerate(footnotes):
                c = ws.cell(row=foot_row + i, column=1, value=note)
                c.font = Font(size=9, color="808080")

            ws.freeze_panes = "A2"
            sheet_summaries.append(f"- **{sheet_title}**：{len(header)} 列 × {len(data_rows)} 行")

        stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        filename = f"table_{_sanitize_filename_part(title)}_{stamp}.xlsx"
        output_dir = Path(_OUTPUT_DIR)
        output_dir.mkdir(parents=True, exist_ok=True)
        output_path = output_dir / filename
        wb.save(str(output_path))

        # Signed download link: The browser clicks directly without JWT header and must use presigned URL for authentication.
        from urllib.parse import quote

        from core.deliverables.download_links import sign_deliverable_query

        download_url = (
            f"/api/assistant/deliverables/{quote(filename)}"
            f"?{sign_deliverable_query(filename)}"
        )

        elapsed = time.time() - t0
        return (
            f"### Excel 表格已导出 ✅\n\n"
            f"**{title}** · {len(parsed)} a Sheet · Time consuming to generate {elapsed:.1f} seconds\n\n"
            + "\n".join(sheet_summaries)
            + f"\n\n"
            f"📎 [Download Excel sheet]({download_url})\n\n"
            f"The tabular data comes verbatim from the model results generated in the conversation without any modifications."
            f"{'Description:' + notes if notes else ''}\n\n"
            f"*This table is generated by AI and is for research reference only and does not constitute investment advice.*"
        )
    except Exception as exc:  # noqa: BLE001
        logger.error("[ExportTableExcelTool] Export failed: %s", exc, exc_info=True)
        return f"### Table export failed\n\n**Execution exception**: {exc}\n\n*Please try again or simplify the table content.*"
