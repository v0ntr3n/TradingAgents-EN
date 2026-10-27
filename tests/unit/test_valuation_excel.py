# -*- coding: utf-8 -*-
"""Valuation Calculation Excel Deliverable Single Measurement (A10-D1)

Coverage: legal spec rendering+Verification, formula grid existence, specification blocking rules (PE negative profit/missing field/zero multiple),
File tampering detection (dead value replacement formula/input value drift), skill entrance round trip.
"""

import json

import pytest
from openpyxl import load_workbook

from core.deliverables.valuation_excel import (
    ValuationSpecError,
    generate_valuation_excel,
    render_valuation_excel,
    validate_spec,
    validate_valuation_excel,
)

# ---- Test specifications (values are demonstration caliber, not investment conclusions)----

VALID_SPEC = {
    "ticker": "601127",
    "stock_name": "Thalys",
    "analysis_date": "2026-09-10",
    "share_count": 15.1,
    "current_price": 130.0,
    "earnings_forecast": {
        "history": [
            {"period": "2024A", "revenue": 1451.8, "net_profit": 59.5},
            {"period": "2025A", "revenue": 1900.0, "net_profit": 85.0},
        ],
        "assumptions": [
            {"period": "2026E", "revenue_growth": 0.30, "net_margin": 0.065,
             "rationale": "Increased volume of new models and ramping up of production capacity"},
            {"period": "2027E", "revenue_growth": 0.20, "net_margin": 0.07,
             "rationale": "Net profit margin continues to improve due to scale effect"},
        ],
    },
    "valuation": {
        "scenarios": [
            {"name": "pessimistic", "method": "PE", "multiple": 12, "eps_period": "2026E",
             "benchmark_basis": "12 times lower end of the industry",
             "applicability_check": {"passed": True, "notes": ["Net profit attributable to parent company is positive"]}},
            {"name": "Neutral", "method": "PE", "multiple": 22, "eps_period": "2026E",
             "benchmark_basis": "The lower edge of the 22-28 times range for peers",
             "applicability_check": {"passed": True, "notes": ["Net profit attributable to parent company is positive"]}},
            {"name": "optimistic", "method": "PE", "multiple": 28, "eps_period": "2027E",
             "benchmark_basis": "The upper edge of the 22-28 times range for peers",
             "applicability_check": {"passed": True, "notes": ["Net profit attributable to parent company is positive"]}},
        ],
        "risk_notes": ["The calculations are scenario assumptions and do not constitute investment advice."],
    },
}


def test_valid_spec_passes_spec_validation():
    assert validate_spec(VALID_SPEC) == []


def test_render_and_validate_roundtrip(tmp_path):
    path = str(tmp_path / "valuation.xlsx")
    out = render_valuation_excel(VALID_SPEC, path)

    import os
    assert os.path.exists(out)

    report = validate_valuation_excel(VALID_SPEC, out)
    assert report["passed"], report["errors"]
    assert report["errors"] == []

    # Python Recalculate expected value (calculate self-consistent numerical anchor point)
    # 2026E: Revenue 1900*1.3=2470, net profit 2470*0.065=160.55, EPS=160.55/15.1
    eps_2026 = 2470.0 * 0.065 / 15.1
    # 2027E: Revenue 2470*1.2=2964, net profit 2964*0.07=207.48, EPS=207.48/15.1
    eps_2027 = 2964.0 * 0.07 / 15.1
    summary = report["summary"]["scenarios"]
    assert summary["pessimistic"]["eps"] == pytest.approx(eps_2026, abs=1e-3)
    assert summary["Neutral"]["target_price"] == pytest.approx(eps_2026 * 22, abs=0.1)
    assert summary["optimistic"]["target_price"] == pytest.approx(eps_2027 * 28, abs=0.1)


def test_forecast_cells_are_formulas(tmp_path):
    path = str(tmp_path / "valuation.xlsx")
    render_valuation_excel(VALID_SPEC, path)

    wb = load_workbook(path)
    ws = wb["Profit deduction"]
    # layout: row 4=2024A, OK 5=2025A, OK 6=2026E, OK 7=2027E
    assert ws["B6"].value == "=B5*(1+C6)"
    assert ws["E6"].value == "=B6*D6"
    assert ws["B7"].value == "=B6*(1+C7)"
    # Assume the grid is the input value (not a formula)
    assert ws["C6"].value == 0.30
    assert ws["D7"].value == 0.07

    ws_v = wb["Valuation scenario"]
    # OK 6=Pessimistic: EPS quoted earnings deduction line 6 (2026E), target price = EPS × multiple
    assert ws_v["E6"].value == "=Profit deduction!E6/$B$2"
    assert ws_v["F6"].value == "=E6*C6"
    assert ws_v["B2"].value == 15.1


def test_pe_with_negative_profit_blocked():
    spec = json.loads(json.dumps(VALID_SPEC))
    spec["earnings_forecast"]["history"].insert(
        0, {"period": "2023A", "revenue": 358.4, "net_profit": -24.5}
    )
    spec["valuation"]["scenarios"][0]["eps_period"] = "2023A"

    errors = validate_spec(spec)
    assert any("PE" in e and "Right" in e for e in errors), errors

    with pytest.raises(ValuationSpecError, match="PE"):
        render_valuation_excel(spec, "unused.xlsx")


def test_missing_required_field_blocked():
    spec = json.loads(json.dumps(VALID_SPEC))
    del spec["valuation"]["scenarios"][0]["benchmark_basis"]

    errors = validate_spec(spec)
    assert any("benchmark_basis" in e for e in errors), errors


def test_zero_multiple_and_margin_blocked():
    spec = json.loads(json.dumps(VALID_SPEC))
    spec["valuation"]["scenarios"][0]["multiple"] = 0
    assert any("multiplier is zero" in e for e in validate_spec(spec))

    spec = json.loads(json.dumps(VALID_SPEC))
    spec["earnings_forecast"]["assumptions"][0]["net_margin"] = 0
    assert any("Net interest rate is zero" in e for e in validate_spec(spec))


def test_tampered_dead_value_detected(tmp_path):
    """Forecast row formulas are replaced with dead values → Check interception"""
    path = str(tmp_path / "valuation.xlsx")
    render_valuation_excel(VALID_SPEC, path)

    wb = load_workbook(path)
    wb["Profit deduction"]["B6"] = 2470.0  # dead value replacement formula
    tampered = str(tmp_path / "tampered_dead.xlsx")
    wb.save(tampered)

    report = validate_valuation_excel(VALID_SPEC, tampered)
    assert not report["passed"]
    assert any("Not a formula" in e for e in report["errors"]), report["errors"]


def test_tampered_input_value_detected(tmp_path):
    """Assume that the input box is changed → Inconsistent with specifications, verification interception"""
    path = str(tmp_path / "valuation.xlsx")
    render_valuation_excel(VALID_SPEC, path)

    wb = load_workbook(path)
    wb["Profit deduction"]["C6"] = 0.50  # The growth assumption has been changed
    tampered = str(tmp_path / "tampered_input.xlsx")
    wb.save(tampered)

    report = validate_valuation_excel(VALID_SPEC, tampered)
    assert not report["passed"]
    assert any("The input value of speed increase is inconsistent with the specification." in e for e in report["errors"]), report["errors"]


def test_skill_entry_roundtrip(tmp_path):
    result_json = generate_valuation_excel(
        json.dumps(VALID_SPEC, ensure_ascii=False),
        output_path=str(tmp_path / "skill_out.xlsx"),
    )
    result = json.loads(result_json)
    assert result["success"], result
    assert result["path"].endswith(".xlsx")
    assert result["summary"]["scenarios"]["Neutral"]["eps"] > 0


def test_skill_entry_invalid_spec():
    bad = json.loads(json.dumps(VALID_SPEC))
    bad["valuation"]["scenarios"][0]["multiple"] = -5
    result = json.loads(generate_valuation_excel(json.dumps(bad)))
    assert not result["success"]
    assert result["stage"] == "spec_validation"
