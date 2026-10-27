# -*- coding: utf-8 -*-
"""Generic table Excel export tool test.

Background: Excel export was once tied to the valuation tool (heavy pipeline, 2-6 minutes).
Conversation data export evaluation pipeline resulted in a 420 second timeout. This tool renders deterministically:
Markdown Original text → openpyxl, Zero LLM calls, completed in seconds.
"""

import re
from pathlib import Path

import pytest
from openpyxl import load_workbook

from core.tools.implementations.assistant_ops.export_table_excel_tool import (
    export_table_excel_tool,
    _parse_cell,
    _parse_markdown_tables,
    _sanitize_sheet_name,
)

_TABLE = (
    "## Monthly Sales Forecast\n"
    "| month | Ask the world M9 | Ask the world M7 |\n"
    "|---|---:|---:|\n"
    "| 2026-09 | 10,500 | 8,000 |\n"
    "| 2026-10 | 11,000 | 8,200 |\n"
)


# ---------- pure function ----------

def test_parse_cell_types():
    assert _parse_cell("10,500") == (10500.0, "#,##0.##")
    assert _parse_cell("-3") == (-3, "#,##0")
    assert _parse_cell("12.5%") == (pytest.approx(0.125), "0.00%")
    assert _parse_cell("3.14") == (3.14, "0.##")
    assert _parse_cell("Ask the world M9")[0] == "Ask the world M9"


def test_parse_markdown_tables_multi_sheet():
    md = (
        "## table A\n| a | b |\n|---|---|\n| 1 | 2 |\n\n"
        "## table B\n| x |\n|---|\n| 9 |\n"
    )
    parsed = _parse_markdown_tables(md)
    assert len(parsed) == 2
    assert parsed[0][0] == "table A"
    assert parsed[0][1] == [["a", "b"], ["1", "2"]]
    assert parsed[1][0] == "table B"


def test_parse_markdown_tables_no_table():
    assert _parse_markdown_tables("plain text without tables") == []


def test_sanitize_sheet_name_dedup():
    used = set()
    n1 = _sanitize_sheet_name("sales table", used)
    n2 = _sanitize_sheet_name("sales table", used)
    assert n1 != n2
    assert _sanitize_sheet_name("a/b\\c?*", set()) == "a_b_c__"


# ---------- Tool integration (real rendering files) ----------

def test_export_creates_excel_file(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    result = export_table_excel_tool.invoke({
        "title": "Cyrus revenue forecast",
        "markdown_tables": _TABLE,
        "notes": "Data source: Dianchedi sales list, ASP takes the median guide price",
    })

    assert "Exported" in result
    assert "/api/assistant/deliverables/table_" in result
    assert "1 a Sheet" in result
    assert "Does not constitute investment advice" in result

    m = re.search(r"deliverables/(table_[^)?]+\.xlsx)", result)
    assert m, f"Download link missing: {result}"
    # The file name in the link is URL encoded and needs to be decoded to check the downloaded file.
    from urllib.parse import unquote

    f = Path("data/deliverables") / unquote(m.group(1))
    assert f.exists()
    # Signature download link: click directly on the browser (without JWT header) to download
    assert "exp=" in result and "sig=" in result

    wb = load_workbook(str(f))
    ws = wb["Monthly Sales Forecast"]
    assert ws.cell(1, 1).value == "month"
    assert ws.cell(2, 2).value == 10500.0
    assert ws.cell(3, 3).value == 8200


def test_export_numbers_verbatim(tmp_path, monkeypatch):
    """Numbers are transmitted verbatim: the values in Excel are consistent with the original Markdown text (not rewritten)."""
    monkeypatch.chdir(tmp_path)
    md = (
        "| indicator | 2026E |\n|---|---:|\n"
        "| Q2 calibration error | 0.06% |\n"
        "| Estimated revenue (billion) | 317.29 |\n"
    )
    result = export_table_excel_tool.invoke({
        "title": "Model calibration", "markdown_tables": md,
    })
    assert "Exported" in result

    m = re.search(r"deliverables/(table_[^)?]+\.xlsx)", result)
    from urllib.parse import unquote

    wb = load_workbook(str(Path("data/deliverables") / unquote(m.group(1))))
    ws = wb.active
    assert ws.cell(2, 2).value == pytest.approx(0.0006)
    assert ws.cell(3, 2).value == 317.29


def test_export_no_table_returns_failure(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    result = export_table_excel_tool.invoke({
        "title": "empty", "markdown_tables": "There is no form here",
    })
    assert "Export failed" in result
    assert "was not parsed in the incoming content" in result


def test_tool_registered_in_registry():
    from core.tools import get_tool_registry
    meta = get_tool_registry().get("export_table_excel_tool")
    assert meta is not None, "export_table_excel_tool Must be automatically registered to the tool registry"
    # light Tier (15s): pure CPU rendering, contrasted with valuation tool heavy (180s)
    assert meta.timeout_tier == "light"
